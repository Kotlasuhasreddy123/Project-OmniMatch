"""Build the professor proposal from its editable Markdown source."""

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


ROOT = Path(__file__).parent
SOURCE = ROOT / "Professor_Startup_Proposal.md"
OUTPUT = ROOT / "OmniMatch Professor Review and Startup Validation.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.9)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
first = True
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line or line == "---":
        i += 1
        continue
    if line.startswith("| "):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
            if not all(set(cell) <= {"-", ":", " "} for cell in cells):
                rows.append(cells)
            i += 1
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = "Table Grid"
        for row in rows:
            cells = table.add_row().cells
            for j, value in enumerate(row):
                cells[j].text = value.replace("**", "").replace("`", "")
        continue
    if line.startswith("# "):
        p = doc.add_paragraph(style="Title" if first else "Heading 1")
        p.add_run(line[2:].replace(":", ""))
        first = False
    elif line.startswith("## "):
        doc.add_heading(line[3:].replace("/", " and "), level=1)
    elif line[0].isdigit() and ". " in line[:4]:
        doc.add_paragraph(line.split(". ", 1)[1].replace("**", "").replace("`", ""), style="List Number")
    elif line.startswith("- "):
        doc.add_paragraph(line[2:], style="List Bullet")
    else:
        doc.add_paragraph(line.replace("**", "").replace("`", ""))
    i += 1

doc.save(OUTPUT)
print(OUTPUT)
