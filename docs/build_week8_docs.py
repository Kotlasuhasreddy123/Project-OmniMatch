"""Build the Week 8 evidence pack and October 20 classroom demo guide."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
REPORT = ROOT / "weekly-reports" / "OmniMatch Week 8 Implementation Milestone Evidence.docx"
GUIDE = ROOT / "team-share" / "OmniMatch October 20 Live Demo Guide.docx"


def shade(cell, fill="E8EEF5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def base_doc(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.60)
    section.bottom_margin = Inches(0.60)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(3)
    for name, size in (("Title", 19), ("Heading 1", 13), ("Heading 2", 10.5)):
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(24, 56, 89)
        style.paragraph_format.space_before = Pt(7 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(3)
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p = doc.add_paragraph(subtitle)
    p.style = "Subtitle"
    p.paragraph_format.space_after = Pt(12)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("OmniMatch  |  Synthetic-data academic prototype  |  15 October 2026")
    return doc


def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def table(doc, headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    for i, h in enumerate(headers):
        tbl.rows[0].cells[i].text = h
        shade(tbl.rows[0].cells[i])
        for run in tbl.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    row_pr = tbl.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row_pr.append(repeat)
    for row in rows:
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    if widths:
        for row in tbl.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return tbl


def make_report():
    doc = base_doc(
        "Week 8 implementation milestone — evidence pack",
        "CPSC-69100 | HealthTech Innovation Labs | Instructor: Dr. Khaled Alzoubi",
    )
    para(doc, "Status as of Thursday, October 15, 2026. This is a team review draft based on code and repeatable local tests. Each team member must verify their own contribution before submission. The exact Week 8 Blackboard deadline was not stated in the syllabus excerpt available to us; confirm it in Blackboard.")
    doc.add_heading("1. What works now", 1)
    bullet(doc, "A local Docker stack starts the Next.js interface, FastAPI service, and Qdrant; all three services were healthy on the tested laptop.")
    bullet(doc, "Consent is recorded before microphone use. Speech audio is handled in memory; Whisper base transcribed a synthetic sentence. Users can edit the transcript before entity extraction.")
    bullet(doc, "A clinician explicitly approves or rejects extracted patient terms. Negated and family-history terms are excluded from the matching request.")
    bullet(doc, "Approved terms route to a background match and an inbox. The demo index contains three invented trials, never live trial IDs; a freshness gate blocks referral for these records.")
    bullet(doc, "The result view shows per-criterion labels, missing lab values, an enrollment-readiness ranking score, and evidence-linked transcript anchors. The score is not a clinical probability.")
    doc.add_heading("2. Verification results", 1)
    table(doc, ["Check", "Observed result", "Meaning"], [
        ("Frontend automated tests", "56 / 56 passed (9 suites)", "UI unit tests pass; not a real-microphone test."),
        ("Backend automated tests", "157 / 159 passed, 1 skipped, 1 failed", "The remaining failure is the unmet 95% NLP target; full suite was rerun."),
        ("Backend code coverage", "75% in the full run", "Below any 95% coverage goal; not a safety measure."),
        ("Live text workflow", "1 relevant fictional match in 14.6 s", "Consent → extraction → approval → asynchronous inbox passed once locally."),
        ("Live synthetic speech", "Correct sample transcription in 6.8 s", "In-memory generated speech, not a physical microphone or broad WER evaluation."),
        ("100-transcript negation corpus", "Precision 77.0%; recall 29.7%; F1 42.9%", "The claimed ≥95% NLP target is NOT met."),
    ], [1.68, 2.55, 2.55])
    para(doc, "These numbers describe different measures. Passing 98% of unit tests does not mean 98% medical accuracy. The benchmark remains red by design until the actual model improves; no result has been relabeled to satisfy the target.")
    doc.add_heading("3. Comparison with existing approaches", 1)
    table(doc, ["Reference", "What it already does", "OmniMatch classroom distinction"], [
        ("AWS HealthScribe", "Clinical transcription, confidence scores, preliminary notes, evidence mapping and human review.", "A narrow research-routing workflow after consent and per-entity approval; it does not replace a clinical note."),
        ("Abridge", "Enterprise ambient documentation and clinician-reviewed, evidence-linked notes.", "Small local simulator with a trial inbox and blocked fictional referrals, not EHR integration."),
        ("ClinicalTrials.gov API v2", "Official structured study records for external retrieval.", "When live mode is enabled, the app indexes records and checks status freshness before any routing. Demo mode uses fiction only."),
    ], [1.22, 2.55, 3.01])
    para(doc, "The distinction is a workflow choice, not a claim that no other product has these features. A practice worth adopting is continuous, stratified evaluation with clinician review, as described by Abridge; our tiny synthetic corpus is not enough for clinical claims.")
    doc.add_heading("4. Remaining work and risks", 1)
    bullet(doc, "Raise negation recall and precision using a separate, independently labeled test set; correct the rules or consider a suitable open model, then rerun the same benchmark without changing the expected labels.")
    bullet(doc, "Test physical laptop microphone permission and speech quality in Chrome or Edge. The automated browser could not use a real microphone; generated speech validated only the backend audio path.")
    bullet(doc, "Fix dependency audit findings (45 reported by npm at install time), add production TLS, encryption at rest, secrets management and formal security review before any real-world use. No HIPAA or clinical compliance claim is supported.")
    bullet(doc, "Exercise more live study-status and referral cases, measure repeated runs under load, and test with diverse speakers and noise. Do not use PHI or contact trial sponsors from the classroom demo.")
    doc.add_heading("5. Week 8 team check-in and next steps", 1)
    para(doc, "The syllabus identifies Week 8 as Implementation Milestone-1 (group check-in, 100 points). It lists the October 20 live demo in Week 9. Confirm the actual Week 8 submission deadline and required form in Blackboard. Use this evidence pack as source material; adapt the final report in your own words and confirm each person's work.")
    table(doc, ["Role in project plan", "Before October 20"], [
        ("Suhas — AI/AWS/matching", "Rehearse local startup, speech result, matching explanation and the honest NLP limitation."),
        ("Rosie — backend/security", "Review endpoint/access claims, audit trail and outstanding security risks; verify her own implemented work."),
        ("Venkata — frontend", "Rehearse consent, transcript correction, approval and inbox navigation; verify his own implemented work."),
    ], [2.23, 4.55])
    doc.add_heading("Sources reviewed (APA-style links)", 1)
    for s in [
        "Amazon Web Services. (n.d.). AWS HealthScribe — AWS AI Service Cards. https://docs.aws.amazon.com/ai/responsible-ai/aws-healthscribe/overview.html",
        "Abridge. (n.d.). Pioneering the science of AI evaluation. https://www.abridge.com/ai/science-ai-evaluation",
        "National Library of Medicine. (n.d.). ClinicalTrials.gov Data API. https://clinicaltrials.gov/data-api/api",
        "OmniMatch local test run. (2026, October 15). Docker, pytest, Jest, and synthetic smoke-test logs (this repository; not externally published).",
    ]:
        para(doc, s)
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(REPORT)


def make_guide():
    doc = base_doc(
        "October 20 live demo — simple team guide",
        "15-minute classroom rehearsal | OmniMatch on one Windows laptop",
    )
    para(doc, "This guide is for a synthetic classroom simulation only. Never speak a real patient's details. A green check in the app means the demo is technically ready, not clinically validated. The presentation date is Tuesday, October 20, 2026.")
    doc.add_heading("1. Before anyone enters the room", 1)
    for line in [
        "Plug in power, connect to a reliable network for first-time model download, and start Docker Desktop 30 minutes early. Once models are cached, the demo runs locally.",
        "In PowerShell, go to Chatgpt\\omnimatch-renewed and run: powershell.exe -NoProfile -ExecutionPolicy Bypass -File .\\START_DEMO.ps1. If already built, add -SkipBuild.",
        "Wait for the green ready message. Open http://localhost:3000. Sign in as clinical_staff with the local-only demo password Staff@12345.",
        "Run .\\RUN_SMOKE_TEST.ps1 and .\\RUN_AUDIO_SMOKE_TEST.ps1 using the same PowerShell options shown above. Both should say PASS. Close other audio apps, allow Chrome/Edge microphone permission, and test your actual laptop microphone once.",
        "Use the fictional case below only. Keep the browser on Home or New Session before the audience arrives; leave Docker Desktop running.",
    ]:
        bullet(doc, line)
    doc.add_heading("2. The 15-minute run of show", 1)
    table(doc, ["Time", "Speaker", "Say / show"], [
        ("0:00–2:30", "Suhas", "Problem, what the prototype is, and its synthetic-only safety boundary."),
        ("2:30–7:30", "Venkata", "Sign in; show consent off, grant consent, record the scripted visit, edit transcript if needed, review entity chips."),
        ("7:30–11:30", "Rosie", "Approve true patient entities, reject errors, submit, open inbox; explain backend background task, access control and trial freshness block."),
        ("11:30–15:00", "Suhas", "Explain criterion labels, unknown lab, heuristic score, what existing products do, honest limitations and next steps."),
    ], [0.9, 1.05, 4.83])
    doc.add_heading("3. Exactly what to do in the app", 1)
    steps = [
        "Open New Session. Point out that Start Recording is unavailable while consent is off. Ask the simulated patient for permission; then click Grant Consent. Only now start the microphone.",
        "Read the patient lines below slowly, with a short pause between sentences. Stop the recording and wait for transcription. If a term is wrong, correct the transcript on the review screen before pressing Proceed.",
        "On the verification screen, approve only current patient facts: type 2 diabetes, metformin, and high blood pressure if detected. Do not approve kidney disease (negated) or stroke (family history). Reject any mistaken chip. Submit verified entities.",
        "Open the new inbox item. Expect one relevant DEMO trial in the current local seed. Show that HbA1c can be UNKNOWN, and that the synthetic trial is BLOCKED from coordinator referral. Do not call it an eligible patient or click a sponsor action.",
        "If results take a few seconds, refresh the inbox. If they do not appear by 60 seconds, use the preflight smoke-test result and explain that this is a prototype limitation rather than inventing a result.",
    ]
    for i, s in enumerate(steps, 1):
        para(doc, f"{i}. {s}")
    doc.add_page_break()
    doc.add_heading("4. Spoken simulation — one person can play both roles", 1)
    para(doc, "Before pressing Grant Consent (not recorded):")
    para(doc, "DOCTOR: “For this fictional classroom demonstration, may I record this visit?”")
    para(doc, "PATIENT: “Yes, I consent.”")
    para(doc, "After consent is granted and Start Recording is clicked, speak these lines clearly:")
    para(doc, "DOCTOR: “Tell me your medical conditions and medicines.”")
    para(doc, "PATIENT: “I have type 2 diabetes. I take metformin. I also have high blood pressure. I do not have kidney disease. My mother had a stroke.”")
    para(doc, "DOCTOR: “Thank you. We will review any possible research study after the visit.”")
    para(doc, "Say the patient sentence at a natural pace. Whisper can mishear accents, noise, or numbers; use the edit step honestly. Do not read out an invented HbA1c number, because the point is to show a missing lab variable.")
    doc.add_heading("5. Short speaking scripts", 1)
    para(doc, "SUHAS — opening: “OmniMatch is our master's project prototype. In a normal visit, trial screening is easy to forget. Our demo listens only after consent, turns the conversation into text, asks a human to approve the facts, and places possible research studies in an inbox for later review. It uses fake data today; it is not for patient care.”")
    para(doc, "VENKATA — interface: “Notice that recording is disabled. After the patient agrees, I grant consent; now the microphone can start. If speech recognition is wrong, I correct the text before continuing. These chips are suggestions. I approve only facts about this patient; family history and negatives do not go into the search.”")
    para(doc, "ROSIE — backend: “The approved terms go to the matching service, which processes the visit after it ends and writes results to an inbox. Access requires a signed-in staff role. We also check whether trial data is current before routing; our fictional demo trial is deliberately blocked from referral. That is a safety gate, not proof of eligibility.”")
    para(doc, "SUHAS — close: “AWS HealthScribe and Abridge already do sophisticated ambient documentation. We are exploring a narrower consent-to-research-inbox workflow, not claiming to be first or clinically superior. Our 100-transcript negation test is well below 95% recall, and we still need better evaluation and production security. The next phase is to improve and independently test the model, then ask our professor for feedback before considering a startup.”")
    doc.add_heading("6. If something fails during the demo", 1)
    table(doc, ["Problem", "Calm response"], [
        ("Microphone denied", "Check browser permission and the consent switch. If blocked, show the in-memory speech PASS output and continue with the text workflow; say the browser microphone was not validated."),
        ("Transcript imperfect", "Use the built-in edit screen, explain human review, and continue. Do not claim flawless transcription."),
        ("No inbox result yet", "Refresh once after 10–20 seconds. If absent, show the prior smoke-test PASS and report the live failure honestly."),
        ("Docker not ready", "Run START_DEMO.ps1 again after Docker Desktop is running; do not present stale screenshots as a live result."),
        ("Question about 95%", "Say: “The current benchmark does not meet it. Test coverage is not model accuracy. Our Week 8 report states the measured precision, recall and F1.”"),
    ], [1.35, 5.43])
    doc.add_heading("7. Team feedback after rehearsal", 1)
    para(doc, "Ask each teammate to try the consent switch, read the patient sentence, correct a transcript, approve one chip, reject one chip, and open the inbox. Record: what worked, any error message, how many seconds it took, and one confusing screen. Bring those comments to the professor. The next technical priorities are higher NLP recall on a held-out set, physical-microphone evaluation, dependency fixes, and production security controls.")
    GUIDE.parent.mkdir(parents=True, exist_ok=True)
    doc.save(GUIDE)


if __name__ == "__main__":
    make_report()
    make_guide()
    print(REPORT)
    print(GUIDE)
